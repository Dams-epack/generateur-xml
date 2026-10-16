#!/usr/bin/env python3
"""Excel -> XML e-Pack Manager (modules produits), passé ensuite dans le script
de la skill completer-produits-xml (complete_xml_products.py) pour une sortie
identique aux profils validés.

  python3 excel_to_xml.py FICHIER.xlsx --list
  python3 excel_to_xml.py FICHIER.xlsx --modules Service,Stickers --out-dir SORTIE
  python3 excel_to_xml.py FICHIER.xlsx --modules all --out-dir SORTIE
"""
import argparse, datetime as dt, os, re, secrets, subprocess, sys, tempfile
from xml.sax.saxutils import escape
from openpyxl import load_workbook
from modules import MODULES

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.join(HERE, "complete_xml_products.py")
TS_RE = re.compile(r"^\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}$")
NUM_RE = re.compile(r"^-?\d+(\.\d+)?$")
BOOLS = {"true":"true","vrai":"true","oui":"true","1":"true","x":"true",
         "false":"false","faux":"false","non":"false","0":"false"}

def cell_text(v, kind, warns, where):
    if v is None: return ""
    if isinstance(v, bool): return "true" if v else "false"
    if isinstance(v, dt.datetime): return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, float) and v.is_integer(): v = int(v)
    if isinstance(v, int) and kind == "id":
        warns.append(f"{where}: Id numérique (zéros initiaux possiblement perdus) -> complété à 18 chiffres")
        return str(v).zfill(18)
    return str(v).strip()

def read_sheet(ws, name, m):
    errs, warns, rows = [], [], []
    hdr = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    idx = {h:i for i,h in enumerate(hdr) if h}
    unknown = [h for h in hdr if h and h not in dict(m["columns"])]
    if unknown: warns.append(f"{name}: colonnes ignorées (inconnues) : {', '.join(unknown)}")
    for t in m["required"]:
        if t not in idx: errs.append(f"{name}: colonne obligatoire absente : {t}")
    if errs: return rows, errs, warns
    for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if all(v is None or str(v).strip()=="" for v in row): continue
        rec = {}
        for tag, kind in m["columns"]:
            where = f"{name} L{r} {tag}"
            raw = row[idx[tag]] if tag in idx and idx[tag] < len(row) else None
            val = cell_text(raw, kind, warns, where)
            if kind == "bool" and val:
                b = BOOLS.get(val.lower())
                if b is None: errs.append(f"{where}: booléen invalide « {val} »"); continue
                val = b
            elif kind == "int" and val:
                if not re.fullmatch(r"-?\d+", val.replace(",", ".").split(".")[0]) or ("." in val.replace(",",".") and float(val.replace(",", "."))%1):
                    errs.append(f"{where}: entier attendu, reçu « {val} »"); continue
                val = str(int(float(val.replace(",", "."))))
            elif kind == "num" and val:
                v2 = val.replace(",", ".")
                if not NUM_RE.match(v2): errs.append(f"{where}: nombre attendu, reçu « {val} »"); continue
                val = v2[:-2] if v2.endswith(".0") else v2
            elif kind == "ts" and val and not TS_RE.match(val):
                errs.append(f"{where}: TS invalide « {val} » (AAAA-MM-JJ HH:MM:SS)"); continue
            elif kind == "id" and val and not val.isdigit():
                errs.append(f"{where}: Id non numérique « {val} »"); continue
            if not val and tag in m["required"]:
                errs.append(f"{where}: valeur obligatoire manquante"); continue
            rec[tag] = val
        rec["_row"] = r
        rows.append(rec)
    return rows, errs, warns

def assign_ids(name, rows, warns, errs):
    now = dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    used, gen = set(), 0
    for rec in rows:
        if rec.get("Id"):
            if rec["Id"] in used: errs.append(f"{name} L{rec['_row']}: Id en double {rec['Id']}")
            used.add(rec["Id"])
    for rec in rows:
        if not rec.get("Id"):
            while True:
                i = str(secrets.randbelow(10**18)).zfill(18)
                if i not in used: break
            used.add(i); rec["Id"] = i; gen += 1
        if not rec.get("TS"): rec["TS"] = now
        if not rec.get("Delete"): rec["Delete"] = "false"
    if gen: warns.append(f"{name}: {gen} Id généré(s) (nouveaux produits)")
    seen = {}
    for rec in rows:
        k = (rec["Name"].casefold(), rec.get("Category","").casefold())
        if k in seen: warns.append(f"{name}: doublon Nom+Catégorie « {rec['Name']} » (L{seen[k]} et L{rec['_row']})")
        else: seen[k] = rec["_row"]

def to_xml(m, rows):
    el = m["element"]
    out = ['<?xml version="1.0" encoding="UTF-8"?>',
           f'<ArrayOf{el} xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance" xmlns:xsd="http://www.w3.org/2001/XMLSchema">']
    for rec in rows:
        out.append(f"  <{el}>")
        for tag, _ in m["columns"]:
            v = rec.get(tag, "")
            if not v and tag in m.get("omit_if_empty", ()): continue
            out.append(f"    <{tag}>{escape(v)}</{tag}>" if v else f"    <{tag}/>")
        out.append(f"  </{el}>")
    out.append(f"</ArrayOf{el}>"); out.append("")
    return "\n".join(out)

def read_params(wb):
    p = {}
    if "Paramètres" in wb.sheetnames:
        for k, v, *_ in wb["Paramètres"].iter_rows(min_row=2, values_only=True):
            if k: p[str(k).strip()] = "" if v is None else str(v).strip()
    return p

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx"); ap.add_argument("--list", action="store_true")
    ap.add_argument("--modules", default=""); ap.add_argument("--out-dir", default="sortie_xml")
    ap.add_argument("--inspect", action="store_true", help="liste onglets, en-têtes et premières lignes")
    ap.add_argument("--sheet", help="onglet source (fichier client libre)")
    ap.add_argument("--header-row", type=int, default=1)
    ap.add_argument("--map", help='JSON {"Name":"Désignation","DateLimit":"=72"} ; "=x" = valeur fixe')
    ap.add_argument("--dltype"); ap.add_argument("--isddm", default="false")
    ap.add_argument("--isddm-exception-category"); ap.add_argument("--isddm-exception-value", default="true")
    a = ap.parse_args()
    wb = load_workbook(a.xlsx, data_only=True)
    if a.inspect:
        for ws in wb.worksheets:
            print(f"\n## Onglet « {ws.title} » ({ws.max_row} lignes x {ws.max_column} colonnes)")
            for i, row in enumerate(ws.iter_rows(min_row=1, max_row=8, values_only=True), 1):
                print(f"  L{i}: {list(row)}")
        return
    if a.map:
        import json
        if len(a.modules.split(",")) != 1 or not a.sheet: sys.exit("--map : un seul module (--modules X) et --sheet requis")
        mod = a.modules.strip(); m0 = MODULES[mod]; mp = json.loads(a.map)
        src = wb[a.sheet]; rows_ = list(src.iter_rows(min_row=a.header_row, values_only=True))
        sh = [str(h).strip() if h is not None else "" for h in rows_[0]]
        for t, c in mp.items():
            if t not in dict(m0["columns"]): sys.exit(f"Balise inconnue dans --map : {t}")
            if not c.startswith("=") and c not in sh: sys.exit(f"Colonne source introuvable : {c} ; disponibles : {sh}")
        from openpyxl import Workbook
        nwb = Workbook(); nws = nwb.active; nws.title = mod
        tags = [t for t, _ in m0["columns"] if t in mp]
        nws.append(tags)
        for r in rows_[1:]:
            if all(v is None or str(v).strip() == "" for v in r): continue
            nws.append([c[1:] if c.startswith("=") else (r[sh.index(c)] if sh.index(c) < len(r) else None) for c in (mp[t] for t in tags)])
        nps = nwb.create_sheet("Paramètres"); nps.append(["Paramètre", "Valeur"])
        for k, v in (("DLType", a.dltype or ""), ("IsDDM", a.isddm), ("IsDDM_exception_category", a.isddm_exception_category or ""), ("IsDDM_exception_value", a.isddm_exception_value)): nps.append([k, v])
        wb = nwb
    filled = {}
    for n in MODULES:
        if n in wb.sheetnames:
            c = sum(1 for r in wb[n].iter_rows(min_row=2, values_only=True)
                    if any(v is not None and str(v).strip() for v in r))
            if c: filled[n] = c
    if a.list:
        for n, c in filled.items(): print(f"{n}: {c} ligne(s)")
        if not filled: print("Aucun module rempli.")
        return
    wanted = list(filled) if a.modules.lower() in ("all", "tous") else [x.strip() for x in a.modules.split(",") if x.strip()]
    if not wanted: sys.exit("Précisez --modules (ou all). Modules remplis : " + ", ".join(filled))
    bad = [w for w in wanted if w not in MODULES]
    if bad: sys.exit(f"Module(s) inconnu(s) : {bad}. Connus : {list(MODULES)}")
    params = read_params(wb)
    os.makedirs(a.out_dir, exist_ok=True)
    rc = 0
    for n in wanted:
        m = MODULES[n]
        print(f"\n=== {n} ===")
        if n not in filled: print("  onglet vide/absent : ignoré"); continue
        rows, errs, warns = read_sheet(wb[n], n, m)
        if not errs: assign_ids(n, rows, warns, errs)
        for w in warns: print("  ⚠", w)
        if errs:
            for e in errs: print("  ✗", e)
            print("  -> module NON généré"); rc = 1; continue
        cmd = [sys.executable, SKILL]
        if m.get("needs_params"):
            if params.get("DLType","") == "":
                print("  ✗ Paramètres!DLType manquant -> module NON généré"); rc = 1; continue
            cmd += ["--dltype", params["DLType"], "--isddm", (params.get("IsDDM") or "false").lower()]
            if params.get("IsDDM_exception_category"):
                cmd += ["--isddm-exception-category", params["IsDDM_exception_category"],
                        "--isddm-exception-value", (params.get("IsDDM_exception_value") or "true").lower()]
        with tempfile.TemporaryDirectory() as td:
            raw = os.path.join(td, m["output"]); open(raw, "w", encoding="utf-8", newline="").write(to_xml(m, rows))
            out = os.path.join(a.out_dir, m["output"])
            r = subprocess.run(cmd + [raw, out], capture_output=True, text=True)
            print("  " + r.stdout.strip().replace("\n", "\n  "))
            if r.returncode: print("  ✗", r.stderr.strip()); rc = 1; continue
            r2 = subprocess.run(cmd + [out, os.path.join(td, "re.xml")], capture_output=True, text=True)
            same = open(out,"rb").read() == open(os.path.join(td,"re.xml"),"rb").read()
            print(f"  idempotence : {'OK' if same else 'ÉCART'} ; {len(rows)} objet(s)")
    sys.exit(rc)

if __name__ == "__main__":
    main()
